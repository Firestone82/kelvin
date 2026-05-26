#!/usr/bin/env python3
"""Generátor obhajovací prezentace: Integrace LLM do systému Kelvin."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- Barevné schéma ----
NAVY = RGBColor(0x14, 0x2A, 0x4F)
BLUE = RGBColor(0x1E, 0x6F, 0xC4)
ACCENT = RGBColor(0x2E, 0xA0, 0x6B)
LIGHT = RGBColor(0xF2, 0xF5, 0xF9)
GRAY = RGBColor(0x5A, 0x66, 0x72)
DARK = RGBColor(0x1B, 0x22, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def box(slide, x, y, w, h):
    from pptx.enum.shapes import MSO_SHAPE
    return slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.0):
    """runs: list of (text, size, bold, color, level) or list of paragraphs each a list of runs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        if isinstance(para, tuple):
            para = [para]
        for (t, size, bold, color, level) in para:
            p.level = level
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "Calibri"
    return tb


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


def header(slide, title, kicker=None):
    bar = box(slide, 0, 0, SW, Inches(1.15))
    fill(bar, NAVY)
    accent = box(slide, 0, Inches(1.15), SW, Inches(0.06))
    fill(accent, ACCENT)
    text(slide, Inches(0.6), Inches(0.18), Inches(12), Inches(0.8),
         [(title, 30, True, WHITE, 0)], anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(slide, Inches(0.62), Inches(0.05), Inches(12), Inches(0.3),
             [(kicker, 12, True, RGBColor(0x9C, 0xC4, 0xEC), 0)])


def bullets(slide, items, x=Inches(0.7), y=Inches(1.6), w=Inches(12), h=Inches(5.5),
            size=18, gap=10):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            label, sub = it
            paras.append([("●  ", size, True, ACCENT, 0), (label, size, True, DARK, 0)])
            if sub:
                paras.append([("     " + sub, size - 3, False, GRAY, 0)])
        else:
            paras.append([("●  ", size, True, ACCENT, 0), (it, size, False, DARK, 0)])
    text(slide, x, y, w, h, paras, space_after=gap, line_spacing=1.05)


# =========================================================
# Slide 1 — Titulní
# =========================================================
s = add_slide()
bg = box(s, 0, 0, SW, SH)
fill(bg, NAVY)
band = box(s, 0, Inches(4.55), SW, Inches(0.08))
fill(band, ACCENT)
text(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.4),
     [("DIPLOMOVÁ PRÁCE  ·  OBHAJOBA", 15, True, RGBColor(0x9C, 0xC4, 0xEC), 0)])
text(s, Inches(0.9), Inches(2.1), Inches(11.6), Inches(2.0),
     [("Integrace velkých jazykových", 44, True, WHITE, 0),
      ("modelů do systému Kelvin", 44, True, WHITE, 0)], line_spacing=1.0)
text(s, Inches(0.9), Inches(4.8), Inches(11.5), Inches(0.6),
     [("AI-asistovaná revize studentských řešení kódu", 22, False, RGBColor(0xC8, 0xD8, 0xEC), 0)])
text(s, Inches(0.9), Inches(6.2), Inches(11.5), Inches(0.8),
     [[("Bc. Pavel Mikula", 18, True, WHITE, 0)],
      [("VŠB – Technická univerzita Ostrava, FEI", 14, False, RGBColor(0x9C, 0xC4, 0xEC), 0)]])
notes(s, "Cca 0:30. Dobrý den, jmenuji se Pavel Mikula. Tématem mé diplomové práce "
         "je integrace velkých jazykových modelů do výukového systému Kelvin – konkrétně "
         "AI-asistovaná revize studentských řešení. Provedu vás motivací, návrhem, "
         "implementací a výsledky.")

# =========================================================
# Slide 2 — Kontext: co je Kelvin
# =========================================================
s = add_slide()
header(s, "Kontext: systém Kelvin", "ÚVOD")
bullets(s, [
    ("Výukový a odevzdávací systém FEI VŠB-TUO", "nasazen na kelvin.cs.vsb.cz, používán napříč kurzy programování"),
    ("Studenti odevzdávají řešení úloh", "automatické vyhodnocení testy + ruční revize kódu vyučujícím"),
    ("Ruční revize je časově náročná", "stovky odevzdání, opakující se chyby, omezená kapacita vyučujících"),
], y=Inches(1.7), gap=14)
# pravý panel - problém
p = box(s, Inches(8.7), Inches(1.9), Inches(4.0), Inches(4.6))
fill(p, LIGHT)
text(s, Inches(9.0), Inches(2.2), Inches(3.5), Inches(4.0),
     [[("Problém", 18, True, NAVY, 0)],
      [("", 6, False, GRAY, 0)],
      [("Kvalitní zpětná vazba ke kódu", 15, False, DARK, 0)],
      [("nestíhá růst počtu studentů.", 15, False, DARK, 0)],
      [("", 8, False, GRAY, 0)],
      [("Cíl práce:", 15, True, ACCENT, 0)],
      [("pomoci vyučujícímu, ne ho", 15, False, DARK, 0)],
      [("nahradit.", 15, False, DARK, 0)]])
notes(s, "Cca 1:00. Kelvin je odevzdávací systém naší fakulty. Studenti sem nahrávají "
         "řešení, ta se automaticky testují a vyučující k nim píše ruční komentáře. "
         "Právě ruční revize je úzké hrdlo – při stovkách odevzdání se opakují stejné chyby. "
         "Mým cílem nebylo vyučujícího nahradit, ale výrazně mu ulehčit.")

# =========================================================
# Slide 3 — Cíle práce
# =========================================================
s = add_slide()
header(s, "Cíle práce", "ZADÁNÍ")
bullets(s, [
    ("Navrhnout integraci LLM do revizního workflow Kelvinu", None),
    ("Automaticky generovat návrhy komentářů ke kódu", "shrnutí řešení + konkrétní připomínky na úrovni řádků"),
    ("Vyučující zůstává v rozhodovací roli", "návrh přijmout, upravit, odmítnout nebo ohodnotit"),
    ("Podpora lokálních i cloudových modelů", "kód studentů nemusí opustit infrastrukturu fakulty"),
    ("Verzování a hodnocení promptů", "možnost měřit a iterovat kvalitu výstupů"),
], y=Inches(1.75), gap=14, size=19)
notes(s, "Cca 1:00. Stanovil jsem pět cílů: navrhnout integraci LLM do stávajícího "
         "workflow, automaticky generovat návrhy komentářů, ponechat vyučujícího v roli "
         "rozhodovatele, podporovat i lokální modely kvůli ochraně dat studentů, a umožnit "
         "verzování promptů kvůli měření kvality. Tyto cíle prolínají celou prací.")

# =========================================================
# Slide 4 — Architektura (data flow)
# =========================================================
s = add_slide()
header(s, "Architektura řešení", "NÁVRH")
from pptx.enum.shapes import MSO_SHAPE

steps = [
    ("1. Odevzdání", "Student nahraje\nřešení úlohy", BLUE),
    ("2. Async fronta", "django-rq job\n(neblokuje pipeline)", BLUE),
    ("3. LLM revize", "OpenAI-kompat. API\nshrnutí + návrhy", ACCENT),
    ("4. Uložení", "SuggestedComment\nstav: pending", BLUE),
    ("5. Vyučující", "přijme / upraví /\nodmítne / ohodnotí", NAVY),
]
n = len(steps)
bw, bh = Inches(2.15), Inches(1.5)
gap = Inches(0.30)
total = bw * n + gap * (n - 1)
startx = (SW - total) / 2
y = Inches(2.5)
for i, (t, d, c) in enumerate(steps):
    x = startx + i * (bw + gap)
    card = box(s, x, y, bw, bh)
    fill(card, c)
    text(s, x + Inches(0.1), y + Inches(0.12), bw - Inches(0.2), bh - Inches(0.2),
         [[(t, 15, True, WHITE, 0)], [("", 4, False, WHITE, 0)],
          [(d, 11.5, False, RGBColor(0xE6, 0xEE, 0xF6), 0)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2, line_spacing=1.0)
    if i < n - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Emu(int(gap)//6),
                                y + bh/2 - Inches(0.12), Inches(0.22), Inches(0.24))
        fill(ar, GRAY)
text(s, Inches(0.7), Inches(4.5), Inches(12), Inches(1.2),
     [[("Klíčové: revize běží asynchronně na pozadí ", 17, False, DARK, 0),
       ("→ neblokuje odevzdání ani testování.", 17, True, ACCENT, 0)],
      [("Výsledky se ukládají jako návrhy ve stavu ", 17, False, DARK, 0),
       ("pending", 17, True, NAVY, 0),
       (" a čekají na vyučujícího.", 17, False, DARK, 0)]],
     space_after=8)
notes(s, "Cca 1:30. Toto je jádro návrhu. Když student odevzdá, do fronty django-rq se "
         "zařadí úloha – běží asynchronně, takže neblokuje odevzdání ani automatické testy. "
         "Job stáhne soubory, zavolá LLM přes OpenAI-kompatibilní API, dostane shrnutí a "
         "návrhy připomínek. Ty se uloží jako objekty SuggestedComment ve stavu pending. "
         "Vyučující je pak v rozhraní zpracuje. Zdůrazním asynchronnost a stav pending – "
         "AI nikdy nepíše přímo studentovi.")

# =========================================================
# Slide 5 — Jak probíhá volání LLM
# =========================================================
s = add_slide()
header(s, "Zpracování v jazykovém modelu", "IMPLEMENTACE")
bullets(s, [
    ("Embedding zdrojových souborů", "podporované jazyky C/C++, Rust, Python, Java; řádky se číslují kvůli přesné referenci"),
    ("Sestavení zprávy pro model", "systémový prompt (z DB) + očíslované soubory + volitelná instrukce k překladu"),
    ("Strukturovaný výstup (JSON)", "response_format=json_object, temperature 0.2 → stabilní, parsovatelný výstup"),
    ("Výsledek", "summary + seznam připomínek: soubor, řádek, závažnost, vysvětlení"),
], y=Inches(1.7), w=Inches(7.6), gap=12, size=17)
# pravý panel: severity + JSON
p = box(s, Inches(8.6), Inches(1.9), Inches(4.1), Inches(4.7))
fill(p, DARK)
text(s, Inches(8.9), Inches(2.1), Inches(3.6), Inches(0.5),
     [("Závažnost připomínek", 15, True, WHITE, 0)])
sev = [("CRITICAL", RGBColor(0xE5, 0x3E, 0x3E)), ("HIGH", RGBColor(0xF2, 0x8B, 0x30)),
       ("MEDIUM", RGBColor(0xE8, 0xC4, 0x4D)), ("LOW", RGBColor(0x4D, 0xB6, 0x6A))]
yy = Inches(2.65)
for label, col in sev:
    chip = box(s, Inches(8.9), yy, Inches(1.7), Inches(0.42))
    fill(chip, col)
    text(s, Inches(8.9), yy, Inches(1.7), Inches(0.42),
         [(label, 12, True, WHITE, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    yy = Emu(int(yy) + int(Inches(0.52)))
text(s, Inches(8.9), Inches(5.0), Inches(3.6), Inches(1.5),
     [[("Výstup mapován na řádky kódu", 12.5, False, RGBColor(0xC8, 0xD8, 0xEC), 0)],
      [("a zobrazen přímo v revizním", 12.5, False, RGBColor(0xC8, 0xD8, 0xEC), 0)],
      [("rozhraní Kelvinu.", 12.5, False, RGBColor(0xC8, 0xD8, 0xEC), 0)]], space_after=2)
notes(s, "Cca 1:15. Detail implementace volání modelu. Stažené soubory očísluji po řádcích, "
         "aby model mohl přesně odkazovat. Sestavím zprávu: systémový prompt z databáze, "
         "zdrojové soubory, a pokud výstup nemá být anglicky, instrukci k překladu. "
         "Vynucuji JSON výstup a nízkou teplotu 0.2 kvůli stabilitě. Model vrátí shrnutí a "
         "připomínky se závažností od LOW po CRITICAL, namapované na konkrétní řádky.")

# =========================================================
# Slide 6 — Provider-agnostic / lokální modely
# =========================================================
s = add_slide()
header(s, "Modely a ochrana dat", "NÁVRHOVÉ ROZHODNUTÍ")
bullets(s, [
    ("OpenAI-kompatibilní rozhraní", "konfigurovatelná base_url → jeden kód, libovolný poskytovatel"),
    ("Lokální modely přes Ollama", "qwen3 / qwen3-coder běží na fakultní infrastruktuře"),
    ("Kód studentů nemusí opustit fakultu", "klíčové pro soukromí a licenční podmínky"),
    ("Více serverů a modelů z konfigurace", "vyučující si vybere server i model pro danou úlohu"),
], y=Inches(1.75), w=Inches(7.6), gap=14, size=18)
p = box(s, Inches(8.6), Inches(1.95), Inches(4.1), Inches(4.5))
fill(p, LIGHT)
text(s, Inches(8.9), Inches(2.25), Inches(3.6), Inches(4.0),
     [[("Konfigurace úlohy", 16, True, NAVY, 0)],
      [("(config.yml)", 12, False, GRAY, 0)],
      [("", 8, False, GRAY, 0)],
      [("async:", 13, True, DARK, 0)],
      [("  llm:", 13, False, DARK, 0)],
      [("    enabled: true", 13, False, ACCENT, 0)],
      [("    language: cs", 13, False, DARK, 0)],
      [("    server_id: local", 13, False, DARK, 0)],
      [("    model: qwen3-coder", 13, False, DARK, 0)],
      [("    prompt_name: default", 13, False, DARK, 0)]], space_after=3, line_spacing=1.05)
notes(s, "Cca 1:00. Důležité návrhové rozhodnutí: použil jsem OpenAI-kompatibilní rozhraní, "
         "takže stačí změnit base_url a stejný kód funguje s OpenAI i s lokálním modelem. "
         "U nás běží lokální Qwen přes Ollama přímo na fakultě – kód studentů tak nemusí "
         "odejít ven, což řeší soukromí. Funkce se zapíná per úlohu v config.yml, vidíte vpravo.")

# =========================================================
# Slide 7 — Workflow vyučujícího + verzování promptů
# =========================================================
s = add_slide()
header(s, "Vyučující a správa promptů", "IMPLEMENTACE")
bullets(s, [
    ("Návrhy zobrazené inline u řádků kódu", "vedle stávajících komentářů, beze změny zvyklostí"),
    ("Akce: přijmout / upravit a přijmout / odmítnout", "přijetí vytvoří reálný komentář a notifikaci studentovi"),
    ("Pouze vyučující dané třídy návrhy ovládá", "kontrola oprávnění na úrovni API"),
    ("Hodnocení návrhů (0–10)", "sběr dat o kvalitě → podklad pro vyhodnocení"),
    ("Verzování promptů v databázi", "name + version, používá se nejnovější verze"),
], y=Inches(1.7), gap=11, size=18)
notes(s, "Cca 1:15. Z pohledu vyučujícího: návrhy se zobrazí přímo u příslušných řádků kódu, "
         "vedle běžných komentářů. Vyučující může návrh přijmout, upravit a přijmout, nebo "
         "odmítnout – přijetí vytvoří skutečný komentář a notifikaci pro studenta. Ovládat "
         "je smí jen vyučující dané třídy. Navíc lze každý návrh ohodnotit 0 až 10, což mi "
         "dává data o kvalitě. Prompty jsou verzované v databázi a vždy se bere nejnovější verze.")

# =========================================================
# Slide 8 — Výsledky / přínos
# =========================================================
s = add_slide()
header(s, "Výsledky a přínos", "ZHODNOCENÍ")
cards = [
    ("Funkční integrace", "End-to-end nasaditelná funkce v reálném Kelvinu", ACCENT),
    ("Human-in-the-loop", "AI navrhuje, vyučující rozhoduje – žádný automat", BLUE),
    ("Provozováno lokálně", "Bez odesílání kódu studentů třetí straně", NAVY),
    ("Měřitelná kvalita", "Hodnocení + verzování promptů pro iteraci", BLUE),
]
cw, ch = Inches(5.9), Inches(2.0)
gx, gy = Inches(0.35), Inches(0.35)
x0, y0 = Inches(0.7), Inches(1.75)
for i, (t, d, c) in enumerate(cards):
    r, col = divmod(i, 2)
    x = x0 + col * (cw + gx)
    y = y0 + r * (ch + gy)
    card = box(s, x, y, cw, ch)
    fill(card, LIGHT)
    bar = box(s, x, y, Inches(0.16), ch)
    fill(bar, c)
    text(s, x + Inches(0.45), y + Inches(0.3), cw - Inches(0.7), ch - Inches(0.5),
         [[(t, 22, True, NAVY, 0)], [("", 6, False, GRAY, 0)],
          [(d, 15, False, DARK, 0)]], space_after=4, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Cca 1:00. Co práce přinesla: funkční end-to-end integraci nasaditelnou v reálném "
         "Kelvinu. Architekturu s člověkem v rozhodovací smyčce – AI jen navrhuje. Možnost "
         "provozu plně lokálně bez odesílání dat ven. A měřitelnost kvality díky hodnocení a "
         "verzování promptů. Všech pět cílů ze zadání bylo naplněno.")

# =========================================================
# Slide 9 — Budoucí práce
# =========================================================
s = add_slide()
header(s, "Možná budoucí rozšíření", "VÝHLED")
bullets(s, [
    ("Rozšíření o další programovací jazyky", "aktuálně C/C++, Rust, Python, Java"),
    ("Kvantitativní studie kvality návrhů", "využití nasbíraných hodnocení napříč kurzy"),
    ("Automatická volba promptu podle typu úlohy", None),
    ("Zohlednění výsledků automatických testů v promptu", "propojení statické revize s dynamickým vyhodnocením"),
], y=Inches(1.8), gap=15, size=19)
notes(s, "Cca 0:40. Kam dál: podpora dalších jazyků, kvantitativní studie kvality nad "
         "nasbíranými hodnoceními, automatická volba promptu podle typu úlohy a propojení "
         "revize s výsledky automatických testů. To už je ale nad rámec této práce.")

# =========================================================
# Slide 10 — Závěr
# =========================================================
s = add_slide()
bg = box(s, 0, 0, SW, SH)
fill(bg, NAVY)
band = box(s, 0, Inches(3.7), SW, Inches(0.08))
fill(band, ACCENT)
text(s, Inches(0.9), Inches(1.3), Inches(11.5), Inches(0.5),
     [("ZÁVĚR", 16, True, RGBColor(0x9C, 0xC4, 0xEC), 0)])
text(s, Inches(0.9), Inches(1.9), Inches(11.6), Inches(1.6),
     [("Velké jazykové modely lze smysluplně", 34, True, WHITE, 0),
      ("integrovat do revize kódu v Kelvinu", 34, True, WHITE, 0)], line_spacing=1.0)
text(s, Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.5),
     [[("Asynchronně, provider-agnosticky, s člověkem v rozhodovací roli.", 19, False, RGBColor(0xC8, 0xD8, 0xEC), 0)],
      [("", 8, False, WHITE, 0)],
      [("Děkuji za pozornost — prostor pro otázky.", 22, True, WHITE, 0)]], space_after=6)
text(s, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.5),
     [("Bc. Pavel Mikula  ·  github.com/Firestone82  ·  systém Kelvin", 13, False, RGBColor(0x9C, 0xC4, 0xEC), 0)])
notes(s, "Cca 0:30. Závěrem: ukázal jsem, že LLM lze do revize kódu v Kelvinu integrovat "
         "smysluplně – asynchronně, nezávisle na poskytovateli modelu a s vyučujícím v "
         "rozhodovací roli. Děkuji za pozornost a jsem připraven na vaše otázky.")

out = "/home/user/kelvin/Prezentace_LLM_Kelvin.pptx"
prs.save(out)
print("Saved:", out, "| slides:", len(prs.slides._sldIdLst))
